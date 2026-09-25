"""
05_build_excel.py
Builds excel/olist_fulfillment_analysis.xlsx from the tables in data/processed/.
Summary tables and KPIs are Excel FORMULAS that calculate from the Clean Order Data sheet,
so Excel checks the Python results independently (see the Data Validation sheet).
"""
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PROCESSED = "data/processed/"
OUTPUT = "excel/olist_fulfillment_analysis.xlsx"

# ---------------------------------------------------------------
# Section 1: Styles (Arial throughout)
# ---------------------------------------------------------------
FONT = "Arial"
TITLE = Font(name=FONT, size=16, bold=True)
SUBTITLE = Font(name=FONT, size=10, italic=True, color="52514E")
HEADER = Font(name=FONT, size=10, bold=True, color="FFFFFF")
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
INPUT = Font(name=FONT, size=10, color="0000FF")   # blue = value typed in from Python
KPI_VALUE = Font(name=FONT, size=24, bold=True, color="1C5CAB")
KPI_LABEL = Font(name=FONT, size=10, bold=True, color="52514E")
HEADER_FILL = PatternFill("solid", fgColor="1C5CAB")
TILE_FILL = PatternFill("solid", fgColor="EEF4FC")
THIN = Side(style="thin", color="D0D0D0")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def write_header(ws, row, headers, start_col=1):
    for i, text in enumerate(headers):
        cell = ws.cell(row=row, column=start_col + i, value=text)
        cell.font = HEADER
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def set_widths(ws, widths):
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


# ---------------------------------------------------------------
# Section 2: Build the Clean Order Data table (option B: key columns only)
# ---------------------------------------------------------------
fact_orders = pd.read_csv(PROCESSED + "fact_orders.csv", parse_dates=["purchase_date"])
dim_customer = pd.read_csv(PROCESSED + "dim_customer.csv", dtype={"customer_zip_code_prefix": str})

rows_before = len(fact_orders)
data = fact_orders.merge(dim_customer[["customer_id", "customer_state"]], on="customer_id", how="left")
assert len(data) == rows_before, "Join changed the number of orders!"

data["purchase_year"] = data["purchase_date"].dt.year
data["purchase_month"] = data["purchase_date"].dt.month

columns = [
    "order_id", "customer_state", "purchase_year", "purchase_month", "order_status",
    "is_delivered", "is_late", "suspicious_date", "total_delivery_time", "transit_time",
    "delivery_variance", "order_value", "freight_cost", "review_score", "review_group",
    "seller_count", "seller_late",
]
data = data[columns]
N = len(data)
LAST = N + 1  # last data row in Excel (row 1 is the header)

# Column letter for each field, e.g. COL["is_late"] -> "G"
COL = {name: get_column_letter(i) for i, name in enumerate(columns, start=1)}


def rng(field):
    """Full-column range on the data sheet, e.g. 'Clean Order Data'!$G$2:$G$99442"""
    return f"'Clean Order Data'!${COL[field]}$2:${COL[field]}${LAST}"


wb = Workbook()
wb.calculation.fullCalcOnLoad = True  # Excel calculates every formula when the file opens

# ---------------------------------------------------------------
# Section 3: README sheet
# ---------------------------------------------------------------
ws = wb.active
ws.title = "README"
readme = [
    ("E-Commerce Fulfillment and Delivery Performance Analysis", TITLE),
    ("Olist Brazilian e-commerce orders, 2016-2018", SUBTITLE),
    ("", BODY),
    ("Business problem", BOLD),
    ("The operations manager at Olist needs to reduce late deliveries because they are resulting in negative "
     "customer review scores. This analysis will help them decide which states and what factors are affecting "
     "late deliveries so they know where to target improvement.", BODY),
    ("", BODY),
    ("Sheets", BOLD),
    ("Dashboard - headline KPIs and charts for the operations manager", BODY),
    ("KPI Definitions - exactly how every KPI is calculated", BODY),
    ("Clean Order Data - one row per order (all 99,441), key columns only", BODY),
    ("Pivot Analysis - summary tables by month, state and lateness, built with Excel formulas", BODY),
    ("Data Validation - Excel's results compared with the Python pipeline's results", BODY),
    ("", BODY),
    ("How this workbook was made", BOLD),
    ("Built by scripts/05_build_excel.py from the tables in data/processed/. To refresh: run scripts 01 to 05 in order.", BODY),
    ("All KPIs and summary tables are live Excel formulas over the Clean Order Data sheet (no pasted results).", BODY),
    ("Blue numbers on the Data Validation sheet are values copied from Python, for comparison.", BODY),
    ("To explore further: select the Clean Order Data table, then Insert > PivotTable.", BODY),
    ("", BODY),
    ("Key definitions", BOLD),
    ("Delivered order: status 'delivered' AND a delivery date. Late: delivered on a later calendar day than promised.", BODY),
    ("Delivery time uses the MEDIAN as the headline (typical order); the average is shown beneath it.", BODY),
    ("Full details: docs/business_questions.md and docs/cleaning_log.md in the project repository.", BODY),
]
for i, (text, font) in enumerate(readme, start=1):
    cell = ws.cell(row=i, column=1, value=text)
    cell.font = font
    cell.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 110

# ---------------------------------------------------------------
# Section 4: KPI Definitions sheet
# ---------------------------------------------------------------
ws = wb.create_sheet("KPI Definitions")
ws["A1"] = "KPI Definitions"
ws["A1"].font = TITLE
ws["A2"] = "Source: docs/business_questions.md. Delivered orders = status 'delivered' AND a delivery date (96,470 orders)."
ws["A2"].font = SUBTITLE
kpis = [
    ("Late delivery rate", "Delivered orders whose delivered calendar day is after the estimated day / delivered orders", "is_delivered = TRUE"),
    ("On-time delivery rate", "Delivered orders delivered on or before the estimated day / delivered orders (= 100% - late rate)", "is_delivered = TRUE"),
    ("Delivery time (median / average)", "Delivered date/time - purchase date/time, in days", "is_delivered = TRUE"),
    ("Delivery variance (days late)", "Delivered calendar day - estimated day. Negative = early, 0 = on the day", "is_delivered = TRUE"),
    ("Transit time", "Delivered date/time - carrier handoff date/time, in days", "valid_transit_time = TRUE (excludes 24 orders)"),
    ("Seller late-handoff rate", "Orders handed to the carrier after the seller's shipping deadline / orders", "seller_count = 1 and a carrier date exists"),
    ("Average review score", "Average of the most recent review score per order (1-5)", "review_score not blank (no review is not counted as 0)"),
    ("Freight percentage", "Total freight / total item price x 100, per order", "orders with items"),
    ("Review group", "Negative (1-2 stars), Neutral (3), Positive (4-5)", "review_score not blank"),
]
write_header(ws, 4, ["KPI", "Definition", "Rows included"])
for r, row in enumerate(kpis, start=5):
    for c, value in enumerate(row, start=1):
        cell = ws.cell(row=r, column=c, value=value)
        cell.font = BOLD if c == 1 else BODY
        cell.alignment = WRAP
        cell.border = BOX
set_widths(ws, [32, 80, 42])

# ---------------------------------------------------------------
# Section 5: Clean Order Data sheet
# ---------------------------------------------------------------
ws = wb.create_sheet("Clean Order Data")
write_header(ws, 1, columns)
for record in data.itertuples(index=False):
    ws.append([None if pd.isna(v) else (v.item() if hasattr(v, "item") else v) for v in record])
for field, fmt in [("total_delivery_time", "0.00"), ("transit_time", "0.00"), ("order_value", "#,##0.00"), ("freight_cost", "#,##0.00")]:
    for cell in ws[COL[field]][1:]:
        cell.number_format = fmt
ws.freeze_panes = "B2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{LAST}"
set_widths(ws, [34, 9, 9, 9, 12, 11, 9, 10, 12, 11, 11, 12, 11, 10, 11, 10, 10])

# ---------------------------------------------------------------
# Section 6: Pivot Analysis sheet (formula-based summary tables)
# ---------------------------------------------------------------
ws = wb.create_sheet("Pivot Analysis")
ws["A1"] = "Pivot Analysis"
ws["A1"].font = TITLE
ws["A2"] = "Every number below is an Excel formula over the Clean Order Data sheet. Delivered orders only."
ws["A2"].font = SUBTITLE

# Table 1: by purchase month
ws["A4"] = "1. Late rate by purchase month"
ws["A4"].font = BOLD
write_header(ws, 5, ["Month", "Year", "Month #", "Delivered orders", "Late orders", "Late rate"])
months = (data[data["is_delivered"]][["purchase_year", "purchase_month"]]
          .drop_duplicates().sort_values(["purchase_year", "purchase_month"]))
month_first_row = 6
for r, (year, month) in enumerate(months.itertuples(index=False), start=month_first_row):
    ws.cell(r, 1, f"{year}-{month:02d}")
    ws.cell(r, 2, int(year))
    ws.cell(r, 3, int(month))
    ws.cell(r, 4, f"=COUNTIFS({rng('purchase_year')},B{r},{rng('purchase_month')},C{r},{rng('is_delivered')},TRUE)")
    ws.cell(r, 5, f"=COUNTIFS({rng('purchase_year')},B{r},{rng('purchase_month')},C{r},{rng('is_late')},TRUE)")
    ws.cell(r, 6, f"=IF(D{r}=0,\"\",E{r}/D{r})").number_format = "0.0%"
month_last_row = month_first_row + len(months) - 1
chart_month_start = month_first_row + int(((months["purchase_year"] == 2016)).sum())  # chart starts Jan 2017
ws.cell(month_last_row + 1, 1, "2016 months have very few orders; the dashboard chart starts at 2017-01.").font = SUBTITLE

# Table 2: by customer state (sorted by number of late orders, computed in Python)
state_top = month_last_row + 3
ws.cell(state_top, 1, "2. Late rate by customer state (sorted by late orders)").font = BOLD
write_header(ws, state_top + 1, ["State", "Delivered orders", "Late orders", "Late rate", "Avg delivery days"])
state_order = (data[data["is_delivered"]].groupby("customer_state")["is_late"].sum()
               .sort_values(ascending=False).index.tolist())
state_first_row = state_top + 2
for r, state in enumerate(state_order, start=state_first_row):
    ws.cell(r, 1, state)
    ws.cell(r, 2, f"=COUNTIFS({rng('customer_state')},A{r},{rng('is_delivered')},TRUE)")
    ws.cell(r, 3, f"=COUNTIFS({rng('customer_state')},A{r},{rng('is_late')},TRUE)")
    ws.cell(r, 4, f"=IF(B{r}=0,\"\",C{r}/B{r})").number_format = "0.0%"
    ws.cell(r, 5, f"=AVERAGEIFS({rng('total_delivery_time')},{rng('customer_state')},A{r},{rng('is_delivered')},TRUE)").number_format = "0.0"
state_last_row = state_first_row + len(state_order) - 1
ws.cell(state_last_row + 1, 1, "Sort order was fixed by Python (most late orders first); values are live formulas.").font = SUBTITLE

# Table 3: review score by how late the order was
late_top = state_last_row + 3
ws.cell(late_top, 1, "3. Review score by delivery variance (delivered orders with a review)").font = BOLD
write_header(ws, late_top + 1, ["Lateness group", "From (days)", "To (days)", "Orders", "Avg review score", "% 1-star"])
groups = [("Early", -1000, -1), ("On the day", 0, 0), ("1-3 days late", 1, 3),
          ("4-7 days late", 4, 7), ("8-14 days late", 8, 14), ("15+ days late", 15, 1000)]
late_first_row = late_top + 2
v, s = rng("delivery_variance"), rng("review_score")
for r, (label, low, high) in enumerate(groups, start=late_first_row):
    ws.cell(r, 1, label)
    ws.cell(r, 2, low)
    ws.cell(r, 3, high)
    ws.cell(r, 4, f"=COUNTIFS({v},\">=\"&B{r},{v},\"<=\"&C{r},{s},\">=1\")")
    ws.cell(r, 5, f"=AVERAGEIFS({s},{v},\">=\"&B{r},{v},\"<=\"&C{r})").number_format = "0.00"
    ws.cell(r, 6, f"=IF(D{r}=0,\"\",COUNTIFS({v},\">=\"&B{r},{v},\"<=\"&C{r},{s},1)/D{r})").number_format = "0.0%"
late_last_row = late_first_row + len(groups) - 1
ws.cell(late_last_row + 1, 1, "-1000 and 1000 stand for 'no limit'.").font = SUBTITLE

for row in ws.iter_rows(min_row=6):
    for cell in row:
        if not cell.font.bold and not cell.font.italic:  # leave headers, titles and notes alone
            cell.font = BODY
set_widths(ws, [30, 16, 12, 16, 16, 12])

# ---------------------------------------------------------------
# Section 7: Data Validation sheet (Excel formulas vs Python values)
# ---------------------------------------------------------------
ws = wb.create_sheet("Data Validation")
ws["A1"] = "Data Validation: Excel vs Python"
ws["A1"].font = TITLE
ws["A2"] = ("Blue = value calculated by the Python pipeline (data/processed/fact_orders.csv), copied in by "
            "scripts/05_build_excel.py. Excel value = live formula over Clean Order Data.")
ws["A2"].font = SUBTITLE
delivered = fact_orders[fact_orders["is_delivered"]]
# Tolerance: counts must match exactly; money within half a cent; rates and averages within 0.000001.
# (Python and Excel add decimals in a different order, so tiny rounding differences are normal.)
checks = [
    ("Orders (rows)", len(fact_orders), f"=COUNTA({rng('order_id')})", "0", 0),
    ("Delivered orders", int(delivered.shape[0]), f"=COUNTIF({rng('is_delivered')},TRUE)", "#,##0", 0),
    ("Late orders", int(fact_orders["is_late"].sum()), f"=COUNTIF({rng('is_late')},TRUE)", "#,##0", 0),
    ("Late rate", float(delivered["is_late"].mean()), f"=COUNTIF({rng('is_late')},TRUE)/COUNTIF({rng('is_delivered')},TRUE)", "0.000%", 0.000001),
    ("Median delivery time (days)", float(delivered["total_delivery_time"].median()), f"=MEDIAN({rng('total_delivery_time')})", "0.0000", 0.000001),
    ("Average delivery time (days)", float(delivered["total_delivery_time"].mean()), f"=AVERAGE({rng('total_delivery_time')})", "0.0000", 0.000001),
    ("Average review score (delivered)", float(delivered["review_score"].mean()), f"=AVERAGEIFS({rng('review_score')},{rng('is_delivered')},TRUE)", "0.0000", 0.000001),
    ("Total order value", float(fact_orders["order_value"].sum()), f"=SUM({rng('order_value')})", "#,##0.00", 0.005),
    ("Total freight cost", float(fact_orders["freight_cost"].sum()), f"=SUM({rng('freight_cost')})", "#,##0.00", 0.005),
    ("Suspicious-date orders", int(fact_orders["suspicious_date"].sum()), f"=COUNTIF({rng('suspicious_date')},TRUE)", "0", 0),
]
write_header(ws, 4, ["Metric", "Python value", "Excel value", "Difference", "Tolerance", "Result"])
for r, (label, python_value, formula, fmt, tolerance) in enumerate(checks, start=5):
    ws.cell(r, 1, label).font = BODY
    cell = ws.cell(r, 2, python_value)
    cell.font = INPUT
    cell.number_format = fmt
    ws.cell(r, 3, formula).number_format = fmt
    ws.cell(r, 4, f"=C{r}-B{r}").number_format = "0.000000"
    ws.cell(r, 5, tolerance).number_format = "0.000000"
    ws.cell(r, 6, f"=IF(ABS(D{r})<=E{r},\"PASS\",\"FAIL\")")
    for c in range(1, 7):
        ws.cell(r, c).border = BOX
        if c > 2:
            ws.cell(r, c).font = BODY
last_check = 4 + len(checks)
ws.cell(last_check + 1, 1, "Tolerance: counts exact; money within half a cent; rates and averages within 0.000001. "
                           "Python and Excel add decimals in a different order, so tiny differences are normal.").font = SUBTITLE
ws.cell(last_check + 3, 1, "Checks passed").font = BOLD
ws.cell(last_check + 3, 2, f"=COUNTIF(F5:F{last_check},\"PASS\")&\" of \"&COUNTA(F5:F{last_check})").font = BOLD
set_widths(ws, [36, 18, 18, 14, 12, 10])

# ---------------------------------------------------------------
# Section 8: Dashboard sheet
# ---------------------------------------------------------------
ws = wb.create_sheet("Dashboard", 1)  # second tab, right after README
ws["B1"] = "Delivery Performance Dashboard"
ws["B1"].font = TITLE
ws["B2"] = "Olist delivered orders, 2016-2018. Every number is a live formula over the Clean Order Data sheet."
ws["B2"].font = SUBTITLE

tiles = [
    ("B", "LATE DELIVERY RATE", f"=COUNTIF({rng('is_late')},TRUE)/COUNTIF({rng('is_delivered')},TRUE)", "0.0%",
     f"=TEXT(COUNTIF({rng('is_late')},TRUE),\"#,##0\")&\" of \"&TEXT(COUNTIF({rng('is_delivered')},TRUE),\"#,##0\")&\" delivered orders\""),
    ("E", "TYPICAL DELIVERY TIME (MEDIAN)", f"=MEDIAN({rng('total_delivery_time')})", "0.0 \"days\"",
     f"=\"Average: \"&TEXT(AVERAGE({rng('total_delivery_time')}),\"0.0\")&\" days\""),
    ("H", "AVERAGE REVIEW SCORE", f"=AVERAGEIFS({rng('review_score')},{rng('is_delivered')},TRUE)", "0.00 \"/ 5\"",
     f"=\"On time: \"&TEXT(AVERAGEIFS({rng('review_score')},{rng('is_late')},FALSE,{rng('is_delivered')},TRUE),\"0.00\")&\"  |  Late: \"&TEXT(AVERAGEIFS({rng('review_score')},{rng('is_late')},TRUE),\"0.00\")"),
]
for col, label, formula, fmt, note in tiles:
    ws[f"{col}4"] = label
    ws[f"{col}4"].font = KPI_LABEL
    ws[f"{col}5"] = formula
    ws[f"{col}5"].font = KPI_VALUE
    ws[f"{col}5"].number_format = fmt
    ws[f"{col}6"] = note
    ws[f"{col}6"].font = SUBTITLE
    for r in (4, 5, 6):
        for offset in range(2):
            ws.cell(r, ws[f"{col}4"].column + offset).fill = TILE_FILL
ws.row_dimensions[5].height = 36
set_widths(ws, [2, 16, 16, 4, 16, 16, 4, 16, 16, 4])

pivot = wb["Pivot Analysis"]


def tidy_chart(chart):
    """Show both axes (newer Excel hides them unless told), use light gridlines, drop the legend."""
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.y_axis.title = None  # the chart title already names the measure; axis titles overlapped the numbers
    chart.varyColors = False   # one colour for every point (Excel otherwise colours each point differently)
    chart.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E5E4E0"))
    chart.legend = None


chart = LineChart()
chart.title = "Late rate by purchase month"
chart.y_axis.title = "Late orders"
chart.y_axis.number_format = "0%"
chart.height, chart.width = 8, 26
chart.add_data(Reference(pivot, min_col=6, min_row=chart_month_start, max_row=month_last_row), titles_from_data=False)
chart.set_categories(Reference(pivot, min_col=1, min_row=chart_month_start, max_row=month_last_row))
tidy_chart(chart)
chart.y_axis.scaling.min = 0
chart.series[0].smooth = False  # straight lines between months: no invented in-between values
chart.series[0].marker.symbol = "circle"
chart.series[0].marker.size = 6
chart.series[0].marker.graphicalProperties = GraphicalProperties(solidFill="2A78D6", ln=LineProperties(solidFill="2A78D6"))
chart.series[0].graphicalProperties.line.solidFill = "2A78D6"
chart.series[0].graphicalProperties.line.width = 22000
ws.add_chart(chart, "B8")

chart = BarChart()
chart.title = "Late orders by customer state (top 10)"
chart.y_axis.title = "Late orders"
chart.height, chart.width = 8, 13
chart.add_data(Reference(pivot, min_col=3, min_row=state_first_row, max_row=state_first_row + 9), titles_from_data=False)
chart.set_categories(Reference(pivot, min_col=1, min_row=state_first_row, max_row=state_first_row + 9))
tidy_chart(chart)
chart.series[0].graphicalProperties.solidFill = "2A78D6"
ws.add_chart(chart, "B25")

chart = BarChart()
chart.title = "Average review score by lateness"
chart.y_axis.title = "Stars (1-5)"
chart.y_axis.scaling.min = 1
chart.y_axis.scaling.max = 5
chart.height, chart.width = 8, 13
chart.add_data(Reference(pivot, min_col=5, min_row=late_first_row, max_row=late_last_row), titles_from_data=False)
chart.set_categories(Reference(pivot, min_col=1, min_row=late_first_row, max_row=late_last_row))
tidy_chart(chart)
chart.y_axis.number_format = "0.0"
chart.series[0].graphicalProperties.solidFill = "2A78D6"
ws.add_chart(chart, "G25")

ws["B42"] = ("Source: Olist public dataset. Late = delivered on a later calendar day than promised. "
             "Definitions: KPI Definitions sheet. Excel vs Python checks: Data Validation sheet.")
ws["B42"].font = SUBTITLE

# ---------------------------------------------------------------
# Section 9: Save
# ---------------------------------------------------------------
os.makedirs("excel", exist_ok=True)
wb.save(OUTPUT)
print("Saved", OUTPUT)
print("Sheets:", wb.sheetnames)
print("Clean Order Data rows:", N, "| columns:", len(columns))
